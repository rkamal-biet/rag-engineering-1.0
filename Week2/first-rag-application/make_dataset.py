"""
make_dataset.py
===============
Builds the sample knowledge base for Task 3 — "Build Your First RAG".

Use case: NovaCart, a FICTIONAL Indian e-commerce company. Every name, number and
policy below is invented for teaching. Nothing refers to a real company.

We deliberately create one file per supported format so the loaders get exercised:

    data/novacart_employee_handbook.pdf   <- PDF  (4 pages -> tests page numbers)
    data/refund_policy.md                 <- Markdown (current 2026 version)
    data/novacart_plus_terms.docx         <- Word
    data/it_security_faq.txt              <- plain text
    data/archive/refund_policy_2025.md    <- OLD version, used only in the
                                             "duplicate document versions" failure demo

Run:  python make_dataset.py
"""

from pathlib import Path

import pymupdf                      # PyMuPDF - used here to WRITE a PDF
from docx import Document as DocxDocument

DATA_DIR = Path(__file__).parent / "data"
ARCHIVE_DIR = DATA_DIR / "archive"


# --------------------------------------------------------------------------------------
# 1. PDF - Employee handbook (one topic per page, so page citations are meaningful)
# --------------------------------------------------------------------------------------
HANDBOOK_PAGES = [
    (
        "NovaCart Employee Handbook 2026 - Section 1: Working at NovaCart",
        """NovaCart is a fictional e-commerce company headquartered in Bengaluru with offices in Pune and Hyderabad.
This handbook applies to all full-time employees from 1 January 2026.

Standard working hours are 9:30 AM to 6:30 PM, Monday to Friday, including a one-hour lunch break.
Core collaboration hours, when every employee must be reachable, are 11:00 AM to 4:00 PM IST.

New employees complete a 6-month probation period. During probation the notice period is 15 days;
after confirmation the notice period is 60 days.

Every new joiner is assigned an onboarding buddy for the first 30 days. The buddy helps with tools,
team introductions and the first production deployment.""",
    ),
    (
        "Section 2: Leave Policy",
        """Every confirmed employee receives 24 days of paid annual leave per calendar year.
Annual leave is credited monthly at the rate of 2 days per month.

Employees receive 12 days of sick leave per year. Sick leave longer than 3 consecutive days
requires a medical certificate uploaded to the HR portal.

Unused annual leave can be carried forward to the next year, up to a maximum of 8 days.
Any balance above 8 days lapses on 31 December and is not encashed.

Parental leave: primary caregivers receive 26 weeks of paid leave; secondary caregivers receive
4 weeks of paid leave, which must be taken within 6 months of the child's arrival.

NovaCart observes 10 fixed public holidays and 2 floating holidays that employees choose themselves.""",
    ),
    (
        "Section 3: Hybrid and Remote Work",
        """NovaCart follows a hybrid model: employees work from the office 3 days per week.
Tuesday and Thursday are mandatory office days for all teams; the third day is chosen by the team.

Employees may work fully remotely for up to 4 weeks per year, with manager approval
requested at least 2 weeks in advance.

Home-office equipment allowance: every employee may claim up to Rs 25,000 once every 3 years
for a chair, desk, monitor or headset. Claims are submitted through the expense portal with receipts.

Company laptops are refreshed every 4 years. Working from a public network is allowed only
when the NovaCart VPN is connected.""",
    ),
    (
        "Section 4: Expenses and Code of Conduct",
        """Business expenses must be submitted within 30 days of the expense date. Claims older than
30 days are rejected automatically by the expense portal.

Daily meal allowance during business travel is Rs 1,500 in metro cities and Rs 1,000 elsewhere.
Domestic flights must be booked in economy class at least 7 days before travel.

Employees must not accept gifts from vendors worth more than Rs 2,000. Any gift above this
limit must be reported to the ethics team at ethics@novacart.example.

Conflicts of interest, such as holding a stake in a supplier, must be declared within 15 days
of joining or of the conflict arising.""",
    ),
]


def build_pdf(path: Path) -> None:
    pdf = pymupdf.open()
    for title, body in HANDBOOK_PAGES:
        page = pdf.new_page(width=595, height=842)            # A4 in points
        page.insert_text((50, 70), title, fontsize=14, fontname="helv")
        rect = pymupdf.Rect(50, 95, 545, 800)
        page.insert_textbox(rect, body, fontsize=10.5, fontname="helv")
    pdf.save(path)


# --------------------------------------------------------------------------------------
# 2. Markdown - Refund policy (current version + an old conflicting version)
# --------------------------------------------------------------------------------------
REFUND_POLICY_2026 = """# NovaCart Refund and Returns Policy

**Version:** 2026-04 (current)
**Applies to:** all orders placed on or after 1 April 2026

## Return windows

| Category | Return window |
|---|---|
| Electronics (phones, laptops, accessories) | 10 days from delivery |
| Fashion and footwear | 30 days from delivery |
| Home and kitchen | 15 days from delivery |
| Books | 7 days from delivery |

Items must be unused, with original tags and packaging.

## Non-returnable items

Innerwear, personal care products, gift cards, digital downloads and customised products
cannot be returned unless they arrive damaged or defective.

## Refund timelines

- UPI and wallet payments: refunded within 48 hours of the returned item passing quality check.
- Credit and debit cards: 5 to 7 business days.
- Cash on delivery orders: refunded to a bank account or NovaCart wallet within 5 business days.

## Damaged or wrong items

Report a damaged, defective or wrong item within 48 hours of delivery with photos through
the app. NovaCart arranges a free pickup and a replacement or full refund.

## Return pickup fee

Return pickups are free for NovaCart Plus members. Other customers pay a Rs 49 pickup fee,
which is waived when the item is damaged or wrong.
"""

REFUND_POLICY_2025 = """# NovaCart Refund and Returns Policy

**Version:** 2025-01 (ARCHIVED - superseded by version 2026-04)

## Return windows

- Electronics: 7 days from delivery
- Fashion and footwear: 15 days from delivery
- Books: 7 days from delivery

## Refund timelines

- UPI payments: 3 business days
- Credit and debit cards: 7 to 10 business days
"""


# --------------------------------------------------------------------------------------
# 3. DOCX - NovaCart Plus membership terms
#    Contains the "cancellation within 14 days ... only if ..." condition from the
#    chapter, so the tiny-chunk failure demo can split it in half.
# --------------------------------------------------------------------------------------
def build_docx(path: Path) -> None:
    doc = DocxDocument()
    doc.add_heading("NovaCart Plus Membership Terms", level=1)
    doc.add_paragraph("Effective date: 1 February 2026. NovaCart Plus is a paid membership programme.")

    doc.add_heading("1. Plans and pricing", level=2)
    doc.add_paragraph(
        "NovaCart Plus costs Rs 999 per year or Rs 129 per month. The yearly plan includes a "
        "30-day free trial for first-time members. Prices include GST."
    )

    doc.add_heading("2. Member benefits", level=2)
    doc.add_paragraph(
        "Plus members get free delivery on every order with no minimum order value, early access "
        "to sale events 24 hours before other customers, 5% extra cashback on NovaCart Pay, and "
        "free return pickups. Non-members get free delivery only on orders above Rs 499."
    )

    doc.add_heading("3. Cancellation and refunds", level=2)
    doc.add_paragraph(
        "Members can cancel the yearly plan within 14 days of purchase for a full refund, but only "
        "if no Plus benefit has been used, such as a free delivery or early sale access. If a benefit "
        "has been used, the membership is cancelled at the end of the current billing period and no "
        "refund is issued. Monthly plans can be cancelled at any time and stop at the end of the month."
    )

    doc.add_heading("4. Auto-renewal", level=2)
    doc.add_paragraph(
        "Memberships renew automatically. NovaCart sends a reminder email and app notification 7 days "
        "before renewal. Auto-renewal can be switched off under Account > Plus > Manage membership."
    )

    doc.add_heading("5. Family sharing", level=2)
    doc.add_paragraph(
        "A yearly Plus member can share benefits with up to 3 family members living at the same "
        "delivery address. Monthly plans do not include family sharing."
    )
    doc.save(path)


# --------------------------------------------------------------------------------------
# 4. TXT - IT security FAQ
# --------------------------------------------------------------------------------------
IT_SECURITY_FAQ = """NovaCart IT Security FAQ (internal, fictional)

Q: How often must I change my password?
A: Passwords must be changed every 90 days. They need at least 14 characters and cannot reuse
any of your last 5 passwords.

Q: Is multi-factor authentication mandatory?
A: Yes. Multi-factor authentication (MFA) is mandatory for email, the VPN, GitHub and every
production system. Use the NovaCart Authenticator app; SMS codes are not allowed.

Q: What should I do with a suspicious email?
A: Do not click links or open attachments. Use the "Report phishing" button in your mail client
or forward the email to security@novacart.example.

Q: My laptop was lost or stolen. What now?
A: Report it to the IT service desk within 24 hours by calling extension 4357 or raising a
ticket in the IT portal. IT will remotely lock and wipe the device.

Q: Can I install software on my company laptop?
A: Only from the NovaCart Software Center. Anything else needs an approved IT ticket.

Q: Can I use personal USB drives?
A: No. USB storage is blocked on company laptops. Share files through the company drive.

Q: Where do I store customer data?
A: Customer personal data may only be stored in approved production databases. It must never be
copied to laptops, spreadsheets or chat tools.
"""


def main() -> None:
    DATA_DIR.mkdir(exist_ok=True)
    ARCHIVE_DIR.mkdir(exist_ok=True)

    build_pdf(DATA_DIR / "novacart_employee_handbook.pdf")
    (DATA_DIR / "refund_policy.md").write_text(REFUND_POLICY_2026, encoding="utf-8")
    build_docx(DATA_DIR / "novacart_plus_terms.docx")
    (DATA_DIR / "it_security_faq.txt").write_text(IT_SECURITY_FAQ, encoding="utf-8")
    (ARCHIVE_DIR / "refund_policy_2025.md").write_text(REFUND_POLICY_2025, encoding="utf-8")

    for f in sorted(DATA_DIR.rglob("*")):
        if f.is_file():
            print(f"created {f.relative_to(DATA_DIR.parent)}  ({f.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
