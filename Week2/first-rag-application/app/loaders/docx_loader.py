"""Word (.docx) loader. Headings are paragraphs too, so they stay in reading order."""

from pathlib import Path
from typing import List

from docx import Document as DocxDocument

from app.core.utils import stable_id
from app.models.document import Document


def load_docx(path: str | Path) -> List[Document]:
    p = Path(path)
    doc = DocxDocument(p)
    paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
    return [
        Document(
            document_id=stable_id(p.name),
            text="\n".join(paragraphs),
            source=p.name,
            metadata={"filename": p.name, "source_type": "docx", "page": None},
        )
    ]
